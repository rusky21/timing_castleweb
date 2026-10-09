import io
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.db.models import Organization, SearchCampaign

class ExcelExporter:
    """Генератор профессиональных отчетов в Excel (.xlsx)"""

    @classmethod
    def generate_campaign_excel(cls, campaign: SearchCampaign, organizations: List[Organization]) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Лиды и аудит"

        # Стили
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")  # Темно-графитовый
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        
        green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        yellow_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
        orange_fill = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color="E5E7EB"),
            right=Side(style='thin', color="E5E7EB"),
            top=Side(style='thin', color="E5E7EB"),
            bottom=Side(style='thin', color="E5E7EB")
        )

        headers = [
            "№",
            "Компания",
            "Категория",
            "Адрес",
            "Рейтинг",
            "Телефон (карты)",
            "Email (с сайта)",
            "Telegram / Соцсети",
            "Сайт",
            "Статус сайта",
            "SSL (HTTPS)",
            "Адаптивность (Mobile)",
            "Метрика / GA",
            "CMS движок",
            "Lead Score",
            "Боль клиента",
            "Решение для продажи",
            "Скрипт звонка менеджера",
            "Ссылка на карты"
        ]

        # Запись заголовков
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        ws.row_dimensions[1].height = 28

        # Запись строк с лидами
        row_num = 2
        for idx, org in enumerate(organizations, start=1):
            audit = org.audit

            primary_phone = org.phones[0] if (org.phones and len(org.phones) > 0) else ""
            email = audit.extra_emails[0] if (audit and audit.extra_emails and len(audit.extra_emails) > 0) else ""
            
            telegram = ""
            if audit and audit.extra_socials:
                telegram = ", ".join(audit.extra_socials[:2])

            status_badge = audit.status_badge if audit else "NO_WEBSITE"
            lead_score = audit.lead_score if audit else 50
            has_ssl_text = "Есть" if (audit and audit.has_ssl) else "НЕТ"
            is_adaptive_text = "Да" if (audit and audit.is_adaptive) else "НЕТ"
            has_analytics_text = "Да" if (audit and audit.has_analytics) else "НЕТ"
            cms = audit.detected_cms if (audit and audit.detected_cms) else "Самопис/HTML"

            pain = audit.pitch_pain if audit else ""
            solution = audit.pitch_solution if audit else ""
            opening = audit.pitch_opening_phrase if audit else ""

            row_data = [
                idx,
                org.name,
                org.category or "",
                org.address or "",
                org.rating,
                primary_phone,
                email,
                telegram,
                org.website or "Нет",
                status_badge,
                has_ssl_text,
                is_adaptive_text,
                has_analytics_text,
                cms,
                lead_score,
                pain,
                solution,
                opening,
                org.card_url or ""
            ]

            ws.append(row_data)

            # Форматирование и подсветка строки
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_num, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # Цветная подсветка статуса
                if col_idx == 10:  # Статус сайта
                    if status_badge in ("NO_SSL", "NO_WEBSITE", "SITE_DOWN"):
                        cell.fill = red_fill
                    elif status_badge == "NOT_RESPONSIVE":
                        cell.fill = orange_fill
                    elif status_badge == "NO_ANALYTICS":
                        cell.fill = yellow_fill
                    elif status_badge == "HTTPS_OK":
                        cell.fill = green_fill

                # Подсветка SSL
                if col_idx == 11 and has_ssl_text == "НЕТ":
                    cell.fill = red_fill

                # Подсветка адаптивности
                if col_idx == 12 and is_adaptive_text == "НЕТ":
                    cell.fill = orange_fill

                # Подсветка метрики
                if col_idx == 13 and has_analytics_text == "НЕТ":
                    cell.fill = yellow_fill

                # Выделение высокого Lead Score
                if col_idx == 15 and lead_score >= 80:
                    cell.font = Font(name="Arial", bold=True, color="991B1B")

            ws.row_dimensions[row_num].height = 24
            row_num += 1

        # Автоподбор ширины колонок
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            # Ограничиваем ширину для текстовых колонок
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
